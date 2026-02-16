# -*- coding: utf-8 -*-
import streamlit as st
import pandas as pd
import math
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
import io
import re

# --- FONCTIONS UTILITAIRES COMMUNES ---
def clean_col_names(df):
    cols = df.columns
    new_cols = []
    for col in cols:
        new_col = str(col).strip().lower(); new_col = re.sub(r"[ ']", '_', new_col); new_col = re.sub(r'[éèê]', 'e', new_col)
        new_col = new_col.replace('fournisseur_principal', 'fournisseur_principal')
        new_cols.append(new_col)
    df.columns = new_cols
    return df

def rechercher_produits(df, references):
    df_recherche = df.copy()
    df_recherche['reference_article'] = df_recherche['reference_article'].astype(str)
    df_recherche['af_reffourniss'] = df_recherche['af_reffourniss'].astype(str)
    df_recherche['reference_mbt'] = df_recherche['reference_mbt'].astype(str)
    references_str = [str(r) for r in references]
    produits_selectionnes = df_recherche[df_recherche['reference_article'].isin(references_str) | df_recherche['af_reffourniss'].isin(references_str) | df_recherche['reference_mbt'].isin(references_str)].copy()
    return produits_selectionnes

# --- FONCTIONS POUR L'ONGLET 1: CALCULATEUR STANDARD ---
def calculer_prix_de_vente_std(produits_df_modifies, marge):
    df_final = produits_df_modifies.copy()
    df_final['prix_d_achat'] = pd.to_numeric(df_final['prix_d_achat'], errors='coerce')
    df_final['dernier_prix_d_achat'] = pd.to_numeric(df_final['dernier_prix_d_achat'], errors='coerce')
    diviseur = 1 - (marge / 100)
    if diviseur <= 0: diviseur = 0.001
    df_final['PV HT (base P.A.)'] = df_final['prix_d_achat'] / diviseur
    df_final['PV HT (base D.P.A.)'] = df_final['dernier_prix_d_achat'] / diviseur
    return df_final

def calculer_totaux_et_comparaison_std(df_final, arrondi, taux_tva=20.0):
    df_final['quantite'] = pd.to_numeric(df_final['quantite'], errors='coerce').fillna(1)
    total_ht_pa = (df_final['PV HT (base P.A.)'] * df_final['quantite']).sum()
    total_ht_dpa = (df_final['PV HT (base D.P.A.)'] * df_final['quantite']).sum()
    if arrondi:
        total_ht_pa = math.ceil(total_ht_pa) - 0.10
        total_ht_dpa = math.ceil(total_ht_dpa) - 0.10
    total_actuel_ttl = 0
    if 'prix_de_vente' in df_final.columns:
        df_final['prix_de_vente'] = pd.to_numeric(df_final['prix_de_vente'], errors='coerce')
        total_actuel_ttl = (df_final['prix_de_vente'] * df_final['quantite']).sum()
    total_actuel_mbt = 0
    if 'prix_de_vente_mbt' in df_final.columns:
        df_final['prix_de_vente_mbt'] = pd.to_numeric(df_final['prix_de_vente_mbt'], errors='coerce')
        total_actuel_mbt = (df_final['prix_de_vente_mbt'] * df_final['quantite']).sum()
    tva_multiplier = 1 + (taux_tva / 100)
    resume = {
        "total_ht_pa": total_ht_pa, "total_ttc_pa": total_ht_pa * tva_multiplier,
        "total_ht_dpa": total_ht_dpa, "total_ttc_dpa": total_ht_dpa * tva_multiplier,
        "total_actuel_ttl": total_actuel_ttl, "total_actuel_mbt": total_actuel_mbt,
        "total_actuel_ttl_ttc": total_actuel_ttl * tva_multiplier,
        "total_actuel_mbt_ttc": total_actuel_mbt * tva_multiplier
    }
    return resume

def exporter_vers_excel_buffer_std(produits_df, resume, entreprise, marge, arrondi):
    wb = Workbook(); ws = wb.active; ws.title = "Calcul Prix de Vente"
    header_font = Font(bold=True, color="FFFFFF"); header_fill = PatternFill(start_color="008000", end_color="008000", fill_type="solid"); summary_font = Font(bold=True)
    ws.merge_cells('A1:B1'); ws['A1'] = "Parametres du Calcul"; ws['A1'].font = summary_font
    ws['A2'] = "Entreprise:"; ws['B2'] = entreprise; ws['A3'] = "Marge appliquee:"; ws['B3'] = f"{marge}%"; ws['A4'] = "Tarif arrondi:"; ws['B4'] = "Oui" if arrondi else "Non"
    start_row = 6
    ws.cell(row=start_row -1, column=1, value="Détail de la Proposition Tarifaire").font = summary_font
    headers = ['Reference article', 'Designation article', 'Quantite', 'Prix d\'achat', 'PV HT unitaire (base P.A.)', 'Dernier Prix d\'achat', 'PV HT unitaire (base D.P.A.)']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_num, value=header_title); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal='center')
    cols_a_afficher = ['reference_article', 'designation_article', 'quantite', 'prix_d_achat', 'PV HT (base P.A.)', 'dernier_prix_d_achat', 'PV HT (base D.P.A.)']
    for row_index, row_data in enumerate(produits_df[cols_a_afficher].itertuples(index=False), start_row + 1):
        for col_num, cell_value in enumerate(row_data, 1):
            cell = ws.cell(row=row_index, column=col_num, value=cell_value)
            if col_num >= 4: cell.number_format = '#,##0.00'
    summary_row_start = start_row + len(produits_df) + 2
    ws.cell(row=summary_row_start, column=4, value="Total HT:").alignment = Alignment(horizontal='right'); ws.cell(row=summary_row_start + 1, column=4, value="Total TTC:").alignment = Alignment(horizontal='right')
    ws.cell(row=summary_row_start, column=6, value="Total HT:").alignment = Alignment(horizontal='right'); ws.cell(row=summary_row_start + 1, column=6, value="Total TTC:").alignment = Alignment(horizontal='right')
    ws.cell(row=summary_row_start, column=5, value=resume['total_ht_pa']).number_format = '#,##0.00'; ws.cell(row=summary_row_start, column=5).font = summary_font
    ws.cell(row=summary_row_start + 1, column=5, value=resume['total_ttc_pa']).number_format = '#,##0.00'; ws.cell(row=summary_row_start + 1, column=5).font = summary_font
    ws.cell(row=summary_row_start, column=7, value=resume['total_ht_dpa']).number_format = '#,##0.00'; ws.cell(row=summary_row_start, column=7).font = summary_font
    ws.cell(row=summary_row_start + 1, column=7, value=resume['total_ttc_dpa']).number_format = '#,##0.00'; ws.cell(row=summary_row_start + 1, column=7).font = summary_font
    comp_start_row = summary_row_start + 4
    ws.cell(row=comp_start_row - 1, column=1, value="Détail Comparatif (basé sur les prix de vente actuels)").font = summary_font
    headers_comp = ['Reference article', 'Designation article', 'Quantite', 'PV Actuel TTL', 'Total Ligne TTL', 'PV Actuel MBT', 'Total Ligne MBT']
    for col_num, header_title in enumerate(headers_comp, 1):
        cell = ws.cell(row=comp_start_row, column=col_num, value=header_title); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal='center')
    cols_comp = ['reference_article', 'designation_article', 'quantite']
    for row_index, row_data in enumerate(produits_df[cols_comp].itertuples(index=False), comp_start_row + 1):
        for col_num, cell_value in enumerate(row_data, 1): ws.cell(row=row_index, column=col_num, value=cell_value)
        qty = produits_df.iloc[row_index - comp_start_row - 1]['quantite']
        pv_ttl = produits_df.iloc[row_index - comp_start_row - 1].get('prix_de_vente', 0); pv_mbt = produits_df.iloc[row_index - comp_start_row - 1].get('prix_de_vente_mbt', 0)
        ws.cell(row=row_index, column=4, value=pv_ttl).number_format = '#,##0.00'; ws.cell(row=row_index, column=5, value=pv_ttl * qty).number_format = '#,##0.00'
        ws.cell(row=row_index, column=6, value=pv_mbt).number_format = '#,##0.00'; ws.cell(row=row_index, column=7, value=pv_mbt * qty).number_format = '#,##0.00'
    comp_summary_row = comp_start_row + len(produits_df) + 1
    ws.cell(row=comp_summary_row, column=4, value="Total Actuel Tetenal:").alignment = Alignment(horizontal='right'); ws.cell(row=comp_summary_row, column=6, value="Total Actuel MB TECH:").alignment = Alignment(horizontal='right')
    ws.cell(row=comp_summary_row, column=5, value=resume['total_actuel_ttl']).number_format = '#,##0.00'; ws.cell(row=comp_summary_row, column=5).font = summary_font
    ws.cell(row=comp_summary_row, column=7, value=resume['total_actuel_mbt']).number_format = '#,##0.00'; ws.cell(row=comp_summary_row, column=7).font = summary_font
    analysis_start_row = comp_summary_row + 3
    headers_analysis = ['Cible de la Proposition', 'Total Proposé (€)', 'Total Actuel (€)', 'Remise (€)', 'Remise (%)']
    ws.cell(row=analysis_start_row - 1, column=1, value="Analyse des Remises vs. Tarif Actuel TETENAL").font = summary_font
    for col_num, header_title in enumerate(headers_analysis, 1):
        cell = ws.cell(row=analysis_start_row, column=col_num, value=header_title); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal='center')
    ttl_data = [("HT (base P.A.)", resume['total_ht_pa'], resume['total_actuel_ttl']), ("TTC (base P.A.)", resume['total_ttc_pa'], resume['total_actuel_ttl_ttc']), ("HT (base D.P.A.)", resume['total_ht_dpa'], resume['total_actuel_ttl']), ("TTC (base D.P.A.)", resume['total_ttc_dpa'], resume['total_actuel_ttl_ttc'])]
    current_row = analysis_start_row + 1
    for label, prop_total, base_total in ttl_data:
        remise_val = base_total - prop_total; remise_pct = (remise_val / base_total) if base_total > 0 else 0
        ws.cell(row=current_row, column=1, value=label); ws.cell(row=current_row, column=2, value=prop_total).number_format = '#,##0.00€'; ws.cell(row=current_row, column=3, value=base_total).number_format = '#,##0.00€'; ws.cell(row=current_row, column=4, value=remise_val).number_format = '#,##0.00€'; ws.cell(row=current_row, column=5, value=remise_pct).number_format = '0.00%'
        current_row += 1
    mbt_start_row = current_row + 2
    ws.cell(row=mbt_start_row - 1, column=1, value="Analyse des Remises vs. Tarif Actuel MB TECH").font = summary_font
    for col_num, header_title in enumerate(headers_analysis, 1):
        cell = ws.cell(row=mbt_start_row, column=col_num, value=header_title); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal='center')
    mbt_data = [("HT (base P.A.)", resume['total_ht_pa'], resume['total_actuel_mbt']), ("TTC (base P.A.)", resume['total_ttc_pa'], resume['total_actuel_mbt_ttc']), ("HT (base D.P.A.)", resume['total_ht_dpa'], resume['total_actuel_mbt']), ("TTC (base D.P.A.)", resume['total_ttc_dpa'], resume['total_actuel_mbt_ttc'])]
    current_row = mbt_start_row + 1
    for label, prop_total, base_total in mbt_data:
        remise_val = base_total - prop_total; remise_pct = (remise_val / base_total) if base_total > 0 else 0
        ws.cell(row=current_row, column=1, value=label); ws.cell(row=current_row, column=2, value=prop_total).number_format = '#,##0.00€'; ws.cell(row=current_row, column=3, value=base_total).number_format = '#,##0.00€'; ws.cell(row=current_row, column=4, value=remise_val).number_format = '#,##0.00€'; ws.cell(row=current_row, column=5, value=remise_pct).number_format = '0.00%'
        current_row += 1
    for col in ws.columns:
        max_length = 0; column_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value: max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = max_length + 2
    buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
    return buffer

# --- FONCTIONS POUR L'ONGLET 2: CALCULATEUR DE PROMOTIONS ---
def calculer_et_analyser_promo(df_promo):
    methode = st.session_state.promo_methode
    df_promo['prix_achat_promo'] = pd.to_numeric(df_promo['prix_achat_promo'], errors='coerce').fillna(0)
    df_promo['prix_de_vente'] = pd.to_numeric(df_promo.get('prix_de_vente'), errors='coerce').fillna(0)
    if methode == "Appliquer une marge":
        marge = st.session_state.promo_marge
        arrondi = st.session_state.promo_arrondi
        diviseur = 1 - (marge / 100)
        if diviseur <= 0: diviseur = 0.001
        df_promo['pv_promo_ht'] = df_promo['prix_achat_promo'] / diviseur
        if arrondi:
            df_promo['pv_promo_ht'] = df_promo['pv_promo_ht'].apply(lambda x: math.ceil(x) - 0.10 if pd.notna(x) else x)
        df_promo['marge_degagee_%'] = marge / 100
    elif methode == "Définir un prix de vente cible":
        df_promo['pv_promo_ht'] = pd.to_numeric(df_promo['pv_promo_ht'], errors='coerce').fillna(0)
        df_promo['marge_degagee_%'] = (1 - (df_promo['prix_achat_promo'] / df_promo['pv_promo_ht']))
        df_promo['marge_degagee_%'] = df_promo['marge_degagee_%'].fillna(0).replace([float('inf'), -float('inf')], 0)
    df_promo['remise_valeur_€'] = df_promo['prix_de_vente'] - df_promo['pv_promo_ht']
    df_promo['remise_%'] = (df_promo['remise_valeur_€'] / df_promo['prix_de_vente'])
    df_promo['remise_%'] = df_promo['remise_%'].fillna(0).replace([float('inf'), -float('inf')], 0)
    return df_promo

def exporter_promo_vers_excel(df_resultats, nom_promo, methode):
    wb = Workbook(); ws = wb.active; ws.title = "Rapport de Promotion"
    header_font = Font(bold=True, color="FFFFFF"); header_fill = PatternFill(start_color="008000", end_color="008000", fill_type="solid"); summary_font = Font(bold=True)
    ws.merge_cells('A1:B1'); ws['A1'] = "Parametres de la Promotion"; ws['A1'].font = summary_font
    ws['A2'] = "Nom de la promotion:"; ws['B2'] = nom_promo
    ws['A3'] = "Methode de calcul:"; ws['B3'] = methode
    start_row = 5
    ws.cell(row=start_row -1, column=1, value="Détail de la Promotion").font = summary_font
    headers = ['Reference article', 'Designation article', 'Quantite', 'Prix Achat Actuel', 'Prix Achat Promo', 'Prix Vente Actuel', 'Prix Vente Promo', 'Marge Dégagée (%)', 'Remise (€)', 'Remise (%)']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_num, value=header_title); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal='center')
    cols_a_afficher = ['reference_article', 'designation_article', 'quantite', 'prix_d_achat', 'prix_achat_promo', 'prix_de_vente', 'pv_promo_ht', 'marge_degagee_%', 'remise_valeur_€', 'remise_%']
    for row_index, row_data in enumerate(df_resultats[cols_a_afficher].itertuples(index=False), start_row + 1):
        for col_num, cell_value in enumerate(row_data, 1):
            cell = ws.cell(row=row_index, column=col_num, value=cell_value)
            if col_num in [4, 5, 6, 7, 9]: cell.number_format = '#,##0.00€'
            if col_num in [8, 10]: cell.number_format = '0.00%'
    for col in ws.columns:
        max_length = 0; column_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value: max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = max_length + 2
    buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
    return buffer

# --- FONCTIONS POUR L'ONGLET 3: COMPARATEUR TTL-MBT ---
def run_comparison(df, filters, analyze_all):
    df_filtered = df.copy()
    if not analyze_all:
        if filters['references']:
            df_filtered = df_filtered[df_filtered['reference_article'].isin(filters['references'])]
        if 'fournisseur_principal' in df_filtered.columns and filters['fournisseur_principal'] != "Tous":
            df_filtered = df_filtered[df_filtered['fournisseur_principal'] == filters['fournisseur_principal']]
        if 'code_famille' in df_filtered.columns and filters['code_famille'] != "Tous":
            df_filtered = df_filtered[df_filtered['code_famille'] == filters['code_famille']]

    for col in ['prix_de_vente', 'prix_de_vente_mbt', 'dernier_prix_d_achat']:
        if col in df_filtered.columns:
            df_filtered[col] = pd.to_numeric(df_filtered[col], errors='coerce')
        else:
            df_filtered[col] = np.nan
            
    df_comp = df_filtered.dropna(subset=['prix_de_vente', 'prix_de_vente_mbt'])
    df_comp = df_comp[(df_comp['prix_de_vente'] > 0) & (df_comp['prix_de_vente_mbt'] > 0)]

    if df_comp.empty:
        return pd.DataFrame(), pd.DataFrame()

    df_comp['ecart_€'] = df_comp['prix_de_vente_mbt'] - df_comp['prix_de_vente']
    df_comp = df_comp[df_comp['ecart_€'] != 0].copy()
    
    df_comp['ecart_%'] = df_comp['ecart_€'] / df_comp['prix_de_vente']
    df_comp['marge_ttl_%'] = (1 - (df_comp['dernier_prix_d_achat'] / df_comp['prix_de_vente']))
    df_comp['marge_mbt_%'] = (1 - (df_comp['dernier_prix_d_achat'] / df_comp['prix_de_vente_mbt']))

    alertes = []
    for _, row in df_comp.iterrows():
        alerte = []
        if 'tarif_vente_-_valeur_remise' in row and pd.notna(row['tarif_vente_-_valeur_remise']) and row['tarif_vente_-_valeur_remise'] != 0:
            alerte.append("Remise TTL")
        if 'remise_mbt' in row and pd.notna(row['remise_mbt']) and "Gamme" in str(row['remise_mbt']):
            alerte.append("Remise MBT")
        alertes.append(", ".join(alerte))
    df_comp['alertes'] = alertes

    for col in ['ecart_%', 'marge_ttl_%', 'marge_mbt_%']:
        df_comp[col] = df_comp[col].replace([np.inf, -np.inf], np.nan)

    df_moins_cher = df_comp[df_comp['ecart_€'] < 0].copy()
    df_plus_cher = df_comp[df_comp['ecart_€'] > 0].copy()

    if analyze_all and 'fournisseur_principal' in df_comp.columns:
        df_moins_cher = df_moins_cher.sort_values(by=['fournisseur_principal', 'reference_article'])
        df_plus_cher = df_plus_cher.sort_values(by=['fournisseur_principal', 'reference_article'])
        
    return df_moins_cher, df_plus_cher

def exporter_comparaison_vers_excel(df_moins_cher, df_plus_cher, filters, analyze_all):
    wb = Workbook(); wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF"); header_fill = PatternFill(start_color="008000", end_color="008000", fill_type="solid"); summary_font = Font(bold=True)

    def setup_sheet(sheet, title, df):
        sheet.merge_cells('A1:B1'); sheet['A1'] = "Parametres de la Comparaison"; sheet['A1'].font = summary_font
        if analyze_all:
            sheet['A2'] = "Mode d'analyse:"; sheet['B2'] = "Analyse Globale de tous les écarts"
        else:
            sheet['A2'] = "Filtres appliqués:"; sheet['B2'] = f"Fournisseur: {filters['fournisseur_principal']}, Famille: {filters['code_famille']}"
        
        base_headers = ['Reference article', 'Designation article', 'PV TTL', 'PV MBT', 'Ecart (€)', 'Ecart (%)', 'Dernier PA', 'Marge TTL (%)', 'Marge MBT (%)', 'Alertes']
        base_cols = ['reference_article', 'designation_article', 'prix_de_vente', 'prix_de_vente_mbt', 'ecart_€', 'ecart_%', 'dernier_prix_d_achat', 'marge_ttl_%', 'marge_mbt_%', 'alertes']
        
        sample_df = df_moins_cher if not df_moins_cher.empty else df_plus_cher
        final_headers, final_cols = base_headers[:], base_cols[:]
        if not sample_df.empty:
            if 'fournisseur_principal' in sample_df.columns:
                final_headers.insert(0, 'Fournisseur Principal'); final_cols.insert(0, 'fournisseur_principal')
            if 'code_famille' in sample_df.columns:
                final_headers.insert(1, 'Code Famille'); final_cols.insert(1, 'code_famille')
        
        sheet.cell(row=4, column=1, value=title).font = summary_font
        for col_num, header_title in enumerate(final_headers, 1):
            cell = sheet.cell(row=5, column=col_num, value=header_title); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal='center')
        
        if df.empty:
            sheet.cell(row=6, column=1, value="Aucun produit ne correspond à cette catégorie.")
        else:
            for row_index, row_data in enumerate(df[final_cols].itertuples(index=False), 6):
                for col_num, cell_value in enumerate(row_data, 1):
                    sheet.cell(row=row_index, column=col_num, value=cell_value)

        for row in sheet.iter_rows(min_row=6, max_row=5 + len(df), min_col=1, max_col=len(final_headers)):
            for cell in row:
                if final_headers[cell.column - 1] in ['PV TTL', 'PV MBT', 'Ecart (€)', 'Dernier PA']:
                    cell.number_format = '#,##0.00€'
                if final_headers[cell.column - 1] in ['Ecart (%)', 'Marge TTL (%)', 'Marge MBT (%)']:
                    cell.number_format = '0.00%'
        
        for col in sheet.columns:
            max_length = 0; column_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value: max_length = max(max_length, len(str(cell.value)))
            sheet.column_dimensions[column_letter].width = max_length + 2

    setup_sheet(wb.create_sheet(title="MBT Moins Cher"), "✅ MBT est Moins Cher que TTL", df_moins_cher)
    setup_sheet(wb.create_sheet(title="MBT Plus Cher"), "❌ MBT est Plus Cher que TTL", df_plus_cher)
    
    buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
    return buffer

# --- FONCTIONS POUR L'ONGLET 4: VÉRIFICATEUR DE MARGES ---
def run_margin_check(df, threshold, filters, analyze_all):
    df_filtered = df.copy()
    if not analyze_all:
        if filters['references']:
            df_filtered = df_filtered[df_filtered['reference_article'].isin(filters['references'])]
        if 'fournisseur_principal' in df_filtered.columns and filters['fournisseur_principal'] != "Tous":
            df_filtered = df_filtered[df_filtered['fournisseur_principal'] == filters['fournisseur_principal']]
        if 'code_famille' in df_filtered.columns and filters['code_famille'] != "Tous":
            df_filtered = df_filtered[df_filtered['code_famille'] == filters['code_famille']]

    for col in ['prix_de_vente', 'prix_de_vente_mbt', 'dernier_prix_d_achat']:
        if col in df_filtered.columns:
            df_filtered[col] = pd.to_numeric(df_filtered[col], errors='coerce')
        else:
            df_filtered[col] = np.nan

    df_check = df_filtered.dropna(subset=['dernier_prix_d_achat'])
    df_check = df_check[df_check['dernier_prix_d_achat'] > 0]
    
    if df_check.empty:
        return pd.DataFrame(), pd.DataFrame()
        
    df_check['marge_ttl_%'] = (1 - (df_check['dernier_prix_d_achat'] / df_check['prix_de_vente'])).replace([np.inf, -np.inf], np.nan)
    df_check['marge_mbt_%'] = (1 - (df_check['dernier_prix_d_achat'] / df_check['prix_de_vente_mbt'])).replace([np.inf, -np.inf], np.nan)
    
    threshold_decimal = threshold / 100
    alerts_ttl = df_check[df_check['marge_ttl_%'] < threshold_decimal].copy()
    alerts_mbt = df_check[df_check['marge_mbt_%'] < threshold_decimal].copy()
    
    if not alerts_ttl.empty:
        alerts_ttl['alerte_remise'] = alerts_ttl.apply(
            lambda row: "Attention remise" if 'tarif_vente_-_valeur_remise' in row and pd.notna(row['tarif_vente_-_valeur_remise']) and row['tarif_vente_-_valeur_remise'] != 0 else "",
            axis=1
        )
    if not alerts_mbt.empty:
        alerts_mbt['alerte_remise'] = alerts_mbt.apply(
            lambda row: "Attention remise" if 'remise_mbt' in row and pd.notna(row['remise_mbt']) and "Gamme" in str(row['remise_mbt']) else "",
            axis=1
        )
        
    return alerts_ttl, alerts_mbt

def exporter_marges_vers_excel(df_ttl, df_mbt, threshold):
    wb = Workbook(); wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF"); header_fill = PatternFill(start_color="008000", end_color="008000", fill_type="solid"); summary_font = Font(bold=True)
    
    def setup_sheet(sheet, title, df, margin_col_name, pv_col_name):
        sheet['A1'] = f"Analyse des Marges - {title}"; sheet['A1'].font = summary_font
        sheet['A2'] = f"Seuil d'alerte: Marge < {threshold}%"
        
        headers = ['Reference article', 'Designation article', 'Dernier Prix d\'Achat', 'Prix de Vente', 'Marge Actuelle (%)', 'Alerte Remise']
        cols_to_display = ['reference_article', 'designation_article', 'dernier_prix_d_achat', pv_col_name, margin_col_name, 'alerte_remise']
        
        for col_num, header_title in enumerate(headers, 1):
            cell = sheet.cell(row=4, column=col_num, value=header_title); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal='center')
        
        if df.empty:
            sheet.cell(row=5, column=1, value="Aucun produit en alerte de marge.")
        else:
            for row_index, row_data in enumerate(df[cols_to_display].itertuples(index=False), 5):
                for col_num, cell_value in enumerate(row_data, 1):
                    sheet.cell(row=row_index, column=col_num, value=cell_value)
        
        for row in sheet.iter_rows(min_row=5, max_row=4 + len(df), min_col=1, max_col=len(headers)):
            for cell in row:
                if headers[cell.column - 1] in ['Dernier Prix d\'Achat', 'Prix de Vente']:
                    cell.number_format = '#,##0.00€'
                if headers[cell.column - 1] == 'Marge Actuelle (%)':
                    cell.number_format = '0.00%'
        
        for col in sheet.columns:
            max_length = 0; column_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value: max_length = max(max_length, len(str(cell.value)))
            sheet.column_dimensions[column_letter].width = max_length + 2

    setup_sheet(wb.create_sheet("Alertes Marge TTL"), "Tetenal", df_ttl, 'marge_ttl_%', 'prix_de_vente')
    setup_sheet(wb.create_sheet("Alertes Marge MBT"), "MB TECH", df_mbt, 'marge_mbt_%', 'prix_de_vente_mbt')
    
    buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
    return buffer

# --- INTERFACE PRINCIPALE DE L'APPLICATION ---
st.set_page_config(page_title="Outil d'Analyse et de Calcul de Prix", layout="wide")
st.title("Outil d'Analyse et de Calcul de Prix")

if 'df_base_clean' not in st.session_state: st.session_state.df_base_clean = None

with st.sidebar:
    uploaded_file = st.file_uploader("1. Chargez votre fichier de base de données Excel", type=["xlsx", "xls"])
    if uploaded_file:
        df_base = pd.read_excel(uploaded_file, header=1)
        st.session_state.df_base_clean = clean_col_names(df_base)
        st.success("Fichier chargé !")

    st.markdown("---")
    if st.button("🔄 Réinitialiser l'application"):
        df_base = st.session_state.get('df_base_clean')
        st.session_state.clear()
        st.session_state.df_base_clean = df_base
        st.rerun()

tab1, tab2, tab3, tab4 = st.tabs(["Calculateur Standard", "Calculateur de Promotions", "Comparateur TTL-MBT", "Vérificateur de Marges"])

with tab1:
    if st.session_state.df_base_clean is None: st.info("Veuillez charger un fichier Excel pour commencer.")
    else:
        # ... (code de l'onglet 1 inchangé)
        df_base_clean = st.session_state.df_base_clean
        with st.expander("⚙️ Étape 1: Configuration et Sélection des Produits", expanded=True):
            st.text_input("Nom du calcul (pour le nom du fichier exporté)", key="std_nom_calcul", placeholder="Ex: Devis Client Dupont")
            col1, col2 = st.columns(2)
            with col1:
                entreprise = st.selectbox("2. Entreprise", ["Tetenal", "MB TECH"], key="std_entreprise"); marge = st.number_input("4. Marge (%)", min_value=0.0, max_value=99.9, value=30.0, step=0.5, format="%.2f", key="std_marge")
            with col2:
                all_references = sorted(pd.concat([df_base_clean['reference_article'].dropna().astype(str), df_base_clean['af_reffourniss'].dropna().astype(str), df_base_clean['reference_mbt'].dropna().astype(str)]).unique())
                references_selectionnees = st.multiselect("3. Produits", options=all_references, key="std_selection"); arrondi = st.checkbox("5. Appliquer l'arrondi final (total en ,90€)", key="std_arrondi")
            if st.button("Afficher les produits pour modification", key="std_btn_afficher"):
                if not references_selectionnees: st.warning("Veuillez selectionner au moins un produit.")
                else:
                    produits = rechercher_produits(df_base_clean, references_selectionnees).copy(); produits['quantite'] = 1
                    st.session_state.std_df_a_modifier = produits
                    if 'std_resume_final' in st.session_state: del st.session_state.std_resume_final
                    if 'std_produits_finaux' in st.session_state: del st.session_state.std_produits_finaux
        if 'std_df_a_modifier' in st.session_state:
            st.header("Étape 2: Modifiez les prix d'achat et les quantités")
            cols_a_editer = ['reference_article', 'designation_article', 'prix_d_achat', 'dernier_prix_d_achat', 'quantite']
            st.session_state.std_df_modifie = st.data_editor(st.session_state.std_df_a_modifier[cols_a_editer], key="std_editor_prix_achat", disabled=['reference_article', 'designation_article'], num_rows="dynamic")
            if st.button("✅ Calculer les Prix de Vente Finaux", key="std_btn_calculer"):
                df_complet_modifie = st.session_state.std_df_a_modifier.copy(); df_complet_modifie.update(st.session_state.std_df_modifie)
                df_avec_pv = calculer_prix_de_vente_std(df_complet_modifie, st.session_state.std_marge)
                resume = calculer_totaux_et_comparaison_std(df_avec_pv, st.session_state.std_arrondi)
                st.session_state.std_produits_finaux = df_avec_pv; st.session_state.std_resume_final = resume
        if 'std_produits_finaux' in st.session_state and 'std_resume_final' in st.session_state:
            st.header("Étape 3: Résultats Finaux")
            st.subheader("Prix de vente unitaires calculés")
            st.dataframe(st.session_state.std_produits_finaux[['reference_article', 'designation_article', 'quantite', 'PV HT (base P.A.)', 'PV HT (base D.P.A.)']])
            st.subheader("Totaux Finaux (tenant compte des quantités)")
            resume_final = st.session_state.std_resume_final; col1, col2 = st.columns(2)
            with col1:
                st.metric(label="Total HT Final (base P.A.)", value=f"{resume_final['total_ht_pa']:.2f} EUR"); st.metric(label="Total TTC Final (base P.A.)", value=f"{resume_final['total_ttc_pa']:.2f} EUR")
            with col2:
                st.metric(label="Total HT Final (base D.P.A.)", value=f"{resume_final['total_ht_dpa']:.2f} EUR"); st.metric(label="Total TTC Final (base D.P.A.)", value=f"{resume_final['total_ttc_dpa']:.2f} EUR")
            st.subheader("Exporter les résultats")
            excel_buffer = exporter_vers_excel_buffer_std(st.session_state.std_produits_finaux, st.session_state.std_resume_final, st.session_state.std_entreprise, st.session_state.std_marge, st.session_state.std_arrondi)
            date_str = datetime.now().strftime("%Y-%m-%d")
            nom_saisi = st.session_state.get("std_nom_calcul", "").strip(); nom_base_fichier = nom_saisi if nom_saisi else "Proposition_Tarifaire"
            nom_base_fichier_safe = re.sub(r'[\\/*?:"<>|]', "", nom_base_fichier); nom_fichier = f"{nom_base_fichier_safe}_{date_str}.xlsx"
            st.download_button(label="Telecharger le rapport standard", data=excel_buffer, file_name=nom_fichier, mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="dl_std")

with tab2:
    # ... (code de l'onglet 2 inchangé)
    if st.session_state.df_base_clean is None: st.info("Veuillez charger un fichier Excel pour commencer.")
    else:
        df_base_clean = st.session_state.df_base_clean
        with st.expander("⚙️ Étape A: Configuration de la Promotion", expanded=True):
            st.text_input("Nom de la promotion (pour le nom du fichier)", key="promo_nom", placeholder="Ex: Promo Printemps 2024")
            all_references = sorted(pd.concat([df_base_clean['reference_article'].dropna().astype(str), df_base_clean['af_reffourniss'].dropna().astype(str), df_base_clean['reference_mbt'].dropna().astype(str)]).unique())
            promo_references = st.multiselect("Sélectionnez les produits pour la promotion", options=all_references, key="promo_selection")
            if st.button("Préparer la promotion", key="promo_btn_preparer"):
                if not promo_references: st.warning("Veuillez sélectionner au moins un produit.")
                else:
                    produits_promo = rechercher_produits(df_base_clean, promo_references).copy()
                    produits_promo['quantite'] = 1
                    produits_promo['prix_achat_promo'] = pd.NA
                    st.session_state.promo_df_a_modifier = produits_promo
                    if 'promo_resultats' in st.session_state: del st.session_state.promo_resultats
        if 'promo_df_a_modifier' in st.session_state:
            st.header("Étape B: Saisir les coûts et quantités")
            cols_a_editer_promo = ['reference_article', 'designation_article', 'prix_d_achat', 'prix_achat_promo', 'quantite']
            st.session_state.promo_df_modifie = st.data_editor(st.session_state.promo_df_a_modifier[cols_a_editer_promo], disabled=['reference_article', 'designation_article', 'prix_d_achat'], num_rows="dynamic", key="promo_editor")
            st.header("Étape C: Choisir la méthode de calcul")
            st.radio("Méthode de calcul du prix de vente promotionnel:", ("Appliquer une marge", "Définir un prix de vente cible"), key="promo_methode")
            if st.session_state.promo_methode == "Appliquer une marge":
                col1, col2 = st.columns(2)
                with col1: st.number_input("Marge promotionnelle (%)", min_value=0.0, max_value=99.9, value=15.0, step=0.5, format="%.2f", key="promo_marge")
                with col2: st.checkbox("Appliquer l'arrondi (,90€)", key="promo_arrondi")
            else:
                st.info("Veuillez saisir les prix de vente cibles dans le tableau ci-dessous.")
                df_pour_cible = st.session_state.promo_df_modifie.copy()
                if 'pv_promo_ht' not in df_pour_cible.columns: df_pour_cible['pv_promo_ht'] = pd.NA
                st.session_state.promo_df_cible = st.data_editor(df_pour_cible, disabled=['reference_article', 'designation_article', 'prix_d_achat', 'prix_achat_promo', 'quantite'], num_rows="dynamic", key="promo_editor_cible")
            if st.button("🚀 Calculer et Analyser la Promotion", key="promo_btn_calculer"):
                df_a_calculer = st.session_state.promo_df_cible if st.session_state.promo_methode == "Définir un prix de vente cible" else st.session_state.promo_df_modifie
                df_complet = pd.merge(df_a_calculer, st.session_state.df_base_clean, on='reference_article', how='left', suffixes=('', '_original'))
                st.session_state.promo_resultats = calculer_et_analyser_promo(df_complet.copy())
        if 'promo_resultats' in st.session_state:
            st.header("Étape D: Résultats de la Promotion")
            df_res = st.session_state.promo_resultats
            st.dataframe(df_res[['reference_article', 'designation_article', 'quantite', 'prix_achat_promo', 'pv_promo_ht', 'marge_degagee_%', 'remise_valeur_€', 'remise_%']])
            st.subheader("Exporter le rapport de promotion")
            excel_buffer_promo = exporter_promo_vers_excel(df_res, st.session_state.promo_nom, st.session_state.promo_methode)
            date_str = datetime.now().strftime("%Y-%m-%d")
            nom_saisi = st.session_state.get("promo_nom", "").strip(); nom_base_fichier = nom_saisi if nom_saisi else "Rapport_Promotion"
            nom_base_fichier_safe = re.sub(r'[\\/*?:"<>|]', "", nom_base_fichier); nom_fichier = f"{nom_base_fichier_safe}_{date_str}.xlsx"
            st.download_button(label="Telecharger le rapport de promotion", data=excel_buffer_promo, file_name=nom_fichier, mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="dl_promo")

with tab3:
    # ... (code de l'onglet 3 inchangé)
    if st.session_state.df_base_clean is None: st.info("Veuillez charger un fichier Excel pour commencer.")
    else:
        df_base_clean = st.session_state.df_base_clean
        with st.expander("⚙️ Configuration du Comparateur", expanded=True):
            analyze_all = st.checkbox("Analyser tous les écarts de prix (classés par Fournisseur Principal)")
            st.write("--- OU (les filtres ci-dessous seront ignorés si la case ci-dessus est cochée) ---")
            col1, col2 = st.columns(2)
            with col1:
                if 'code_famille' in df_base_clean.columns:
                    familles = ["Tous"] + sorted(df_base_clean['code_famille'].dropna().unique().tolist())
                    sel_famille = st.selectbox("Filtrer par Code Famille", options=familles, key="comp_code_famille")
                else:
                    st.warning("Colonne 'Code Famille' introuvable."); st.session_state.comp_code_famille = "Tous"
            with col2:
                if 'fournisseur_principal' in df_base_clean.columns:
                    fournisseurs = ["Tous"] + sorted(df_base_clean['fournisseur_principal'].dropna().unique().tolist())
                    sel_fournisseur = st.selectbox("Filtrer par Fournisseur Principal", options=fournisseurs, key="comp_fournisseur_principal")
                else:
                    st.warning("Colonne 'Fournisseur principal' introuvable."); st.session_state.comp_fournisseur_principal = "Tous"
            all_references_comp = sorted(pd.concat([df_base_clean['reference_article'].dropna().astype(str), df_base_clean['af_reffourniss'].dropna().astype(str), df_base_clean['reference_mbt'].dropna().astype(str)]).unique())
            sel_references = st.multiselect("Filtrer par Références", options=all_references_comp, key="comp_selection")
            if st.button("🚀 Lancer la Comparaison", key="comp_btn_lancer"):
                filters = {'references': st.session_state.get('comp_selection', []), 'fournisseur_principal': st.session_state.get('comp_fournisseur_principal', 'Tous'), 'code_famille': st.session_state.get('comp_code_famille', 'Tous'),}
                moins_cher, plus_cher = run_comparison(df_base_clean, filters, analyze_all)
                st.session_state.comp_moins_cher = moins_cher; st.session_state.comp_plus_cher = plus_cher
                st.session_state.comp_filters = filters; st.session_state.comp_analyze_all = analyze_all
        if 'comp_moins_cher' in st.session_state:
            st.header("Résultats de la Comparaison")
            cols_to_show = ['reference_article', 'designation_article', 'prix_de_vente', 'prix_de_vente_mbt', 'ecart_€', 'ecart_%', 'dernier_prix_d_achat', 'marge_ttl_%', 'marge_mbt_%', 'alertes']
            if 'fournisseur_principal' in st.session_state.df_base_clean.columns: cols_to_show.insert(0, 'fournisseur_principal')
            if 'code_famille' in st.session_state.df_base_clean.columns: cols_to_show.insert(1, 'code_famille')
            if not st.session_state.comp_moins_cher.empty:
                st.subheader("✅ MBT est Moins Cher que TTL"); st.dataframe(st.session_state.comp_moins_cher[cols_to_show])
            if not st.session_state.comp_plus_cher.empty:
                st.subheader("❌ MBT est Plus Cher que TTL"); st.dataframe(st.session_state.comp_plus_cher[cols_to_show])
            if st.session_state.comp_moins_cher.empty and st.session_state.comp_plus_cher.empty:
                st.info("Aucun écart de prix trouvé pour la sélection actuelle.")
            st.subheader("Exporter le rapport de comparaison")
            excel_buffer_comp = exporter_comparaison_vers_excel(st.session_state.comp_moins_cher, st.session_state.comp_plus_cher, st.session_state.comp_filters, st.session_state.comp_analyze_all)
            date_str = datetime.now().strftime("%Y-%m-%d"); nom_fichier = f"Comparaison_TTL_MBT_{date_str}.xlsx"
            st.download_button(label="Telecharger le rapport de comparaison", data=excel_buffer_comp, file_name=nom_fichier, mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="dl_comp")

with tab4:
    if st.session_state.df_base_clean is None: st.info("Veuillez charger un fichier Excel pour commencer.")
    else:
        df_base_clean = st.session_state.df_base_clean
        with st.expander("⚙️ Configuration du Vérificateur de Marges", expanded=True):
            seuil_marge = st.number_input("Seuil de marge minimum (%)", min_value=0.1, max_value=100.0, value=20.0, step=1.0, key="margin_threshold")
            analyze_all_margins = st.checkbox("Analyser tous les produits")
            st.write("--- OU ---")
            col1, col2 = st.columns(2)
            with col1:
                if 'code_famille' in df_base_clean.columns:
                    familles = ["Tous"] + sorted(df_base_clean['code_famille'].dropna().unique().tolist())
                    sel_famille_margin = st.selectbox("Filtrer par Code Famille", options=familles, key="margin_famille")
                else:
                    st.warning("Colonne 'Code Famille' introuvable."); st.session_state.margin_famille = "Tous"
            with col2:
                if 'fournisseur_principal' in df_base_clean.columns:
                    fournisseurs = ["Tous"] + sorted(df_base_clean['fournisseur_principal'].dropna().unique().tolist())
                    sel_fournisseur_margin = st.selectbox("Filtrer par Fournisseur Principal", options=fournisseurs, key="margin_fournisseur")
                else:
                    st.warning("Colonne 'Fournisseur principal' introuvable."); st.session_state.margin_fournisseur = "Tous"
            all_references_margin = sorted(pd.concat([df_base_clean['reference_article'].dropna().astype(str), df_base_clean['af_reffourniss'].dropna().astype(str), df_base_clean['reference_mbt'].dropna().astype(str)]).unique())
            sel_references_margin = st.multiselect("Filtrer par Références", options=all_references_margin, key="margin_selection")
            if st.button("🚀 Lancer la Vérification des Marges", key="margin_btn_lancer"):
                filters = {'references': st.session_state.get('margin_selection', []), 'fournisseur_principal': st.session_state.get('margin_fournisseur', 'Tous'), 'code_famille': st.session_state.get('margin_famille', 'Tous'),}
                alerts_ttl, alerts_mbt = run_margin_check(df_base_clean, st.session_state.margin_threshold, filters, analyze_all_margins)
                st.session_state.margin_alerts_ttl = alerts_ttl; st.session_state.margin_alerts_mbt = alerts_mbt
        if 'margin_alerts_ttl' in st.session_state:
            st.header("Résultats de la Vérification")
            if not st.session_state.margin_alerts_ttl.empty:
                st.subheader("📉 Alertes de Marge Insuffisante - TTL")
                st.dataframe(st.session_state.margin_alerts_ttl[['reference_article', 'designation_article', 'dernier_prix_d_achat', 'prix_de_vente', 'marge_ttl_%', 'alerte_remise']])
            else:
                st.info("Aucune alerte de marge trouvée pour TTL avec le seuil et les filtres actuels.")
            if not st.session_state.margin_alerts_mbt.empty:
                st.subheader("📉 Alertes de Marge Insuffisante - MBT")
                st.dataframe(st.session_state.margin_alerts_mbt[['reference_article', 'designation_article', 'dernier_prix_d_achat', 'prix_de_vente_mbt', 'marge_mbt_%', 'alerte_remise']])
            else:
                st.info("Aucune alerte de marge trouvée pour MBT avec le seuil et les filtres actuels.")
            st.subheader("Exporter le rapport des marges")
            excel_buffer_margin = exporter_marges_vers_excel(st.session_state.margin_alerts_ttl, st.session_state.margin_alerts_mbt, st.session_state.margin_threshold)
            date_str = datetime.now().strftime("%Y-%m-%d"); nom_fichier = f"Rapport_Marges_{date_str}.xlsx"
            st.download_button(label="Telecharger le rapport des marges", data=excel_buffer_margin, file_name=nom_fichier, mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="dl_margin")
